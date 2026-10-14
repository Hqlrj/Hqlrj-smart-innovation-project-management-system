# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_cn(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size=16 if level == 1 else (14 if level == 2 else 12),
            bold=True,
            name="微软雅黑",
        )
    return p


def add_para(text, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_qa(q, a):
    p = doc.add_paragraph()
    rq = p.add_run("Q：" + q)
    set_run_font(rq, size=11, bold=True, color=(0x1A, 0x56, 0xDB))
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.3

    p2 = doc.add_paragraph()
    ra = p2.add_run("A：" + a)
    set_run_font(ra, size=10.5)
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.line_spacing = 1.35
    p2.paragraph_format.left_indent = Cm(0.3)


# ========== Title ==========
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("智能创新创业项目管理系统")
set_run_font(r, size=20, bold=True)

title2 = doc.add_paragraph()
title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = title2.add_run("面试准备手册（自我介绍 · 技术栈 · 面试题）")
set_run_font(r2, size=14, bold=True, color=(0x33, 0x33, 0x33))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("基于个人项目整理 · 含 Redis / 微服务扩展知识点 · 100+ 题")
set_run_font(rs, size=10, color=(0x66, 0x66, 0x66))

# ========== 一 ==========
add_heading_cn("一、面试自我介绍", 1)

add_heading_cn("1.1 完整版（约 1.5 分钟）", 2)
add_para(
    "面试官你好，我叫【姓名】，是【学校/专业】的【应届生/在校生】。"
    "我近期独立完成了一个前后端分离的「智能创新创业项目管理系统」。"
    "前端采用 Vue 3 + TypeScript + Vite + Element Plus，使用 Pinia 管理登录态，"
    "Axios 统一封装请求与 JWT 鉴权，并用 ECharts 做项目数据可视化，docx-preview 实现计划书在线预览。"
    "后端采用 Java 21 + Spring Boot 3 + MyBatis + MySQL，完成了用户角色管理、项目申报与多级审批、"
    "获奖管理、公告与站内通知、文件上传等核心功能。"
    "项目亮点是接入 DeepSeek 大模型：通过 Apache POI 解析 Word 计划书，按章节拆分后异步调用大模型做 AI 评审，"
    "前端轮询展示进度；同时还做了站内 AI 对话助手。"
    "权限上支持系统管理员、项目管理员、项目负责人、指导老师、评委老师、审批人等多角色，"
    "前端路由菜单控制与后端 JWT 拦截器双重校验。"
    "通过这个项目，我对前后端联调、鉴权、文件处理、业务状态流转和 AI 接口对接都有较完整的实践经验。"
    "希望能加入贵司，在【后端/前端/全栈】方向继续成长，谢谢。"
)

add_heading_cn("1.2 极简版（约 30 秒）", 2)
add_para(
    "我做过一个 Vue3 + Spring Boot 的高校创新创业项目管理系统，覆盖项目申报审批、多角色权限、JWT 登录鉴权，"
    "以及 DeepSeek 计划书 AI 评审与智能对话。"
    "技术栈是 Vue3/TypeScript/Element Plus/Pinia + Spring Boot/MyBatis/MySQL/JWT。希望应聘【岗位】。"
)

add_heading_cn("1.3 分岗位侧重点", 2)
add_para(
    "【前端岗】多提：Composition API、路由守卫、Pinia、Axios 拦截器、Element Plus、ECharts、Word 预览、角色菜单权限。"
)
add_para(
    "【后端岗】多提：分层架构、JWT 拦截器、MyBatis/PageHelper、文件上传、POI、异步 AI 评审、统一异常、事务与状态机。"
)
add_para("【全栈岗】用完整版，前后端都点到，突出联调与业务闭环。")

# ========== 二 ==========
add_heading_cn("二、项目技术栈清单（含补充）", 1)
add_para(
    "说明：标「已用」是项目真实依赖；标「建议掌握」是面试常考、可结合项目谈优化的技术。",
    size=10,
)

add_heading_cn("2.1 整体架构（已用）", 2)
for i in [
    "前后端分离架构",
    "RESTful API（统一响应 Result：code / msg / data）",
    "开发环境：Vite 代理 /api → 后端 8080",
    "生产环境：Nginx 静态托管前端 + 反向代理后端",
    "分层：Controller → Service → Mapper（MyBatis XML）",
]:
    add_para("• " + i, size=10.5)

add_heading_cn("2.2 前端技术（已用）", 2)
for i in [
    "Vue 3.2（Composition API / <script setup>）",
    "TypeScript 4.7",
    "Vite 3（构建与热更新、路径别名 @、开发代理）",
    "Vue Router 4（history 模式、路由懒加载、全局守卫）",
    "Pinia 2（登录态：token→sessionStorage，userInfo→localStorage）",
    "Element Plus 2 + @element-plus/icons-vue（中文语言包）",
    "Axios（请求/响应拦截器、Bearer JWT、401 跳转）",
    "ECharts 5（项目一览图表）",
    "docx-preview-sync + JSZip（Word 计划书前端预览）",
    "ESLint + Prettier + vue-tsc（代码规范与类型检查）",
    "Vitest / Cypress（工程脚手架具备单测与 E2E 能力）",
    "npm 包管理",
]:
    add_para("• " + i, size=10.5)

add_heading_cn("2.3 后端技术（已用）", 2)
for i in [
    "Java 21",
    "Spring Boot 3.2（spring-boot-starter-web，内嵌 Tomcat）",
    "MyBatis 3 + Mapper XML（下划线转驼峰）",
    "PageHelper（分页：total + list）",
    "MySQL 8（库名 bsgl，InnoDB，utf8mb4）",
    "HikariCP（Spring Boot 默认连接池）",
    "JWT / jjwt 0.12（自定义 JwtInterceptor，非 Spring Security）",
    "MD5 密码校验（自定义 Md5Util；面试可谈升级 BCrypt）",
    "Lombok",
    "Apache POI 5.2（poi / poi-ooxml / poi-scratchpad，解析 doc/docx）",
    "OkHttp 4 + Gson（调用 DeepSeek HTTP API）",
    "Jackson（日期格式化等）",
    "文件上传 MultipartFile（头像 / 计划书 / 获奖证明，UUID 重命名）",
    "CORS 跨域配置、静态资源映射 /uploads/**",
    "@ControllerAdvice 全局异常处理",
    "Maven + Maven Wrapper",
    "DeepSeek Chat Completions（AI 评审 + AI 对话；用户可配 API Key）",
    "异步任务思路（AI 评审进度轮询；CompletableFuture / 异步处理）",
    "SHA-256 文件指纹校验（计划书完整性）",
]:
    add_para("• " + i, size=10.5)

add_heading_cn("2.4 业务模块（体现你“做过”）", 2)
for i in [
    "用户注册 / 登录 / 找回密码 / 头像上传 / 改密",
    "角色管理与多角色权限（菜单 + 路由 + 接口）",
    "项目申报（成员、指导老师、计划书）",
    "项目状态机：draft → submitted → approved/rejected → settled",
    "获奖登记与证明审批",
    "公告发布与站内通知（审批结果、公告推送）",
    "Word 计划书上传、在线预览、模板下载",
    "DeepSeek AI 计划书评审（拆段、异步、进度、报告）",
    "站内 AI 对话助手",
    "ECharts 项目/获奖数据一览",
]:
    add_para("• " + i, size=10.5)

add_heading_cn("2.5 面试建议补充掌握（项目暂未用，但强烈建议会说）", 2)
add_para(
    "这些可以写成「我项目目前是单体 + MySQL；如果优化我会引入……」，显得有成长空间。",
    size=10,
)
for i in [
    "Redis：缓存、验证码、未读数、JWT 黑名单、限流、AI 进度缓存",
    "Spring Security：替代手写拦截器，统一认证授权与密码加密",
    "BCrypt / 加盐哈希：替代 MD5",
    "消息队列（RabbitMQ / Kafka）：审批通知、AI 评审异步解耦、削峰",
    "微服务（Spring Cloud）：Gateway、Nacos、OpenFeign、Sentinel",
    "Docker / docker-compose：容器化部署",
    "对象存储 OSS / MinIO：替代本地 uploads 目录",
    "WebSocket / SSE：替代轮询，推送通知与 AI 进度",
    "Swagger / Knife4j / OpenAPI：接口文档",
    "MyBatis-Plus：简化 CRUD（可选了解）",
    "单元测试 JUnit + Mockito；接口测试 Postman / Apifox",
    "Linux 基础命令、Nginx 配置、Git 协作流程",
]:
    add_para("• " + i, size=10.5)

add_heading_cn("2.6 简历技术栈一行（可直接复制）", 2)
add_para(
    "Vue3 / TypeScript / Vite / Element Plus / Pinia / Axios / ECharts / "
    "Spring Boot 3 / Java 21 / MyBatis / MySQL / JWT / Apache POI / DeepSeek"
    "（了解 Redis、微服务、MQ）"
)

# ========== 三 Redis ==========
add_heading_cn("三、Redis 知识点（面试速通）", 1)
add_para(
    "项目现状：当前未使用 Redis。话术：「验证码、未读通知、热点查询如果做优化，我会引入 Redis。」"
)

add_heading_cn("3.1 是什么", 2)
add_para(
    "基于内存的高性能 Key-Value 数据库，常用于缓存、会话、分布式锁、排行榜、限流、轻量消息队列。"
)

add_heading_cn("3.2 五种常用数据类型", 2)
for i in [
    "String：验证码、缓存 JSON、计数器",
    "Hash：用户对象分字段存储",
    "List：最新消息列表、简单队列",
    "Set：去重、共同关注",
    "ZSet（有序集合）：排行榜、延时任务",
]:
    add_para("• " + i)

add_heading_cn("3.3 经典问题", 2)
for i in [
    "缓存穿透：查不存在的数据直打 DB → 布隆过滤器 / 缓存空值",
    "缓存击穿：热点 Key 过期瞬间大量请求 → 互斥锁重建 / 逻辑过期",
    "缓存雪崩：大量 Key 同时过期 → 过期时间加随机、多级缓存、集群高可用",
    "与 MySQL 一致性：常用「先更新 DB，再删除缓存」；或延时双删",
    "持久化：RDB 快照、AOF 日志；生产常混合使用",
    "过期删除：定期删除 + 惰性删除；内存满时淘汰策略如 allkeys-lru",
]:
    add_para("• " + i)

add_heading_cn("3.4 结合本项目可落地的改造点", 2)
for i in [
    "1）找回密码验证码：Redis 存 code:phone，TTL 5 分钟",
    "2）退出登录 / 强制下线：JWT 黑名单放 Redis",
    "3）未读通知数：高频接口用 Redis 计数",
    "4）AI 评审进度：进度数据放 Redis，减少轮询打库",
    "5）DeepSeek 调用限流：按用户/API Key 滑动窗口限流",
]:
    add_para(i)

# ========== 四 微服务 ==========
add_heading_cn("四、微服务知识点（面试速通）", 1)
add_para(
    "项目现状：单体 Spring Boot。话术：「业务清晰、部署简单；量上来或团队变大再拆服务。」"
)

add_heading_cn("4.1 单体 vs 微服务", 2)
add_para("单体：一个应用一个包，开发部署简单，但一挂全挂、扩展粒度粗。")
add_para("微服务：按业务拆多个独立服务，可独立部署扩容，但引入分布式复杂度。")

add_heading_cn("4.2 常见组件（Spring Cloud 体系）", 2)
for i in [
    "注册中心 / 配置中心：Nacos（或 Eureka + Config）",
    "网关：Spring Cloud Gateway（统一入口、鉴权、路由、限流）",
    "远程调用：OpenFeign",
    "负载均衡：Spring Cloud LoadBalancer",
    "熔断降级限流：Sentinel / Resilience4j",
    "链路追踪：SkyWalking / Zipkin",
    "消息队列：RabbitMQ / Kafka（异步解耦）",
]:
    add_para("• " + i)

add_heading_cn("4.3 若拆分本项目（面试加分）", 2)
add_para("API Gateway（JWT 鉴权）")
add_para("  ├─ 用户服务（登录 / 用户 / 角色）")
add_para("  ├─ 项目服务（申报 / 审批 / 计划书）")
add_para("  ├─ 通知服务（公告 / 站内信）")
add_para("  ├─ AI 服务（DeepSeek 评审 / 对话）")
add_para("  └─ 文件服务（上传预览，可对接 OSS）")
add_para(
    "拆分原则：按业务边界拆，优先拆变化快、压力大的模块（AI、文件）；注意分布式事务与最终一致性。"
)

# ========== 五 面试题 ==========
add_heading_cn("五、面试题大全（含参考答案）", 1)

add_heading_cn("5.1 项目综合题", 2)
for q, a in [
    (
        "请介绍一下你的项目。",
        "这是一个面向高校的创新创业项目管理系统，前后端分离。核心流程包括用户注册审批、项目申报（成员/导师/计划书）、"
        "提交与审批驳回、入驻；还有获奖管理、公告、站内通知。亮点是接入 DeepSeek，对计划书做 AI 评审，并提供 AI 对话助手。",
    ),
    (
        "你在项目中负责什么？遇到的最大难点是什么？",
        "（按真实情况说）难点可举：① JWT + 多角色权限前后端配合；② Word 解析与前端预览兼容 doc/docx；"
        "③ AI 评审异步任务与进度轮询、失败重试。说明问题、方案、结果。",
    ),
    (
        "项目状态是怎么流转的？",
        "项目：draft→submitted→approved/rejected→settled；获奖：pending→approved/rejected；"
        "用户：待审批→已审批→已停用。审批与公告会触发站内通知。",
    ),
    (
        "为什么选择前后端分离？",
        "职责清晰：前端专注交互，后端专注业务与数据；便于并行开发、独立部署和联调；通过统一 JSON 接口通信。",
    ),
    (
        "统一响应格式怎么设计的？",
        "Result<T>：code（1成功/0失败）、msg、data。分页用 total + list。前端 Axios 拦截器成功时直接解包 data。",
    ),
    (
        "文件上传如何实现？有什么不足？",
        "MultipartFile 存本地 uploads，UUID 重命名，静态资源映射访问。不足：单机、难扩容、备份麻烦。"
        "改进：OSS/MinIO + CDN，库中只存 URL。",
    ),
    (
        "密码如何存储？如何改进？",
        "当前 MD5 校验。生产应使用 BCrypt 等慢哈希并加盐，且配合 Spring Security 密码编码器；可加登录失败次数限制防爆破。",
    ),
    (
        "AI 评审链路怎么做的？",
        "上传计划书→POI 解析→按章节拆分→创建评审任务→异步调用 DeepSeek→结果写入章节/报告表→前端轮询任务进度展示。"
        "用户可配置自己的 API Key。",
    ),
    (
        "通知为什么用轮询而不是 WebSocket？",
        "实现简单，满足当前规模。改进可用 WebSocket/SSE 推送未读数与 AI 进度，降低无效请求。",
    ),
    (
        "前后端如何联调与部署？",
        "开发：Vite 代理 /api。生产：前端 npm run build 出 dist，Nginx 托管并 try_files 支持 history；"
        "后端 jar，Nginx /api 反代到 8080。",
    ),
    (
        "如何保证接口安全？",
        "JWT 鉴权、角色校验、MyBatis 预编译防注入、文件类型与大小限制、敏感配置不入库明文（可改进）、HTTPS（生产）。",
    ),
    (
        "如果并发变高，你怎么优化？",
        "接口缓存（Redis）、DB 索引与慢 SQL、连接池调优、AI 与通知异步化（MQ）、静态资源 CDN、必要时读写分离或拆服务。",
    ),
    (
        "角色权限如何设计？",
        "前端 permission.ts 按角色控制菜单与路由；后端 JWT 解析 role，接口内再校验（如谁能发公告）。"
        "强调：前端权限只提升体验，安全以后端为准。",
    ),
    (
        "计划书预览怎么实现的？",
        "后端提供文件流/下载接口；前端用 docx-preview-sync 渲染 docx；后端 POI 也用于解析内容给 AI。"
        "旧版 doc 用 poi-scratchpad 处理。",
    ),
    (
        "项目有哪些表？关系大概怎样？",
        "user、role、project、project_member、project_advisor、project_plan、award、notice、notification，"
        "以及 AI 相关 ai_review_task/section/report、ai_conversation/message 等。项目与成员/计划书/获奖多为 1 对多。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.2 前端面试题", 2)
for q, a in [
    (
        "Vue2 和 Vue3 有什么区别？",
        "Vue3 推荐 Composition API，逻辑复用更好；Proxy 响应式；更好的 TypeScript 支持；打包更小、性能更好；"
        "新增 Teleport、Fragments、Suspense 等。",
    ),
    (
        "什么是 Composition API？你项目怎么用？",
        "用 setup / <script setup> 按功能组织代码，而不是按 options 拆分。页面里用 ref/reactive/computed/生命周期组合业务逻辑。",
    ),
    (
        "v-if 和 v-show 区别？",
        "v-if 条件渲染（销毁/重建），适合不常切换；v-show 切 display，适合频繁切换。",
    ),
    (
        "computed 和 watch 区别？",
        "computed 有缓存、依赖收集，适合派生值；watch 监听变化执行副作用（请求、同步存储等）。",
    ),
    (
        "Vue Router 导航守卫做了什么？",
        "beforeEach：恢复登录态；未登录跳登录；已登录校验角色路由权限，无权限提示并回首页。",
    ),
    (
        "为什么用 Pinia 而不是 Vuex？",
        "Pinia 更轻量、API 简洁、天生 TS 友好，是 Vue 官方推荐的下一代状态管理。",
    ),
    (
        "Axios 拦截器你做了哪些事？",
        "请求拦截：除登录注册外附加 Authorization Bearer token。响应拦截：解包 Result；业务失败提示；401 清登录态跳转。",
    ),
    (
        "如何做权限控制？",
        "roleMenuMap / roleRouteMap 控制菜单与路由；后端仍要校验。前端权限不能当安全边界。",
    ),
    (
        "key 在 v-for 里有什么用？",
        "帮助 Virtual DOM diff 复用节点，应用唯一稳定 id，避免用 index 导致状态错乱。",
    ),
    (
        "什么是路由懒加载？",
        "() => import('...') 按路由拆包，加快首屏。项目路由表已采用。",
    ),
    (
        "Element Plus 如何做表单校验？",
        "el-form + :rules + FormInstance.validate，配合 FormRules 类型。",
    ),
    (
        "ECharts 如何避免内存泄漏？",
        "组件卸载时 chart.dispose()，窗口 resize 监听要移除。",
    ),
    (
        "浏览器本地存储有哪些？你们怎么用？",
        "localStorage 持久、sessionStorage 会话级、Cookie。项目：token 用 sessionStorage，userInfo 用 localStorage。",
    ),
    (
        "跨域是什么？你们怎么解决？",
        "浏览器同源策略限制。开发用 Vite 代理；后端也配了 CORS；生产用 Nginx 同域反代。",
    ),
    (
        "说一下 Vue 响应式原理（简要）",
        "Vue3 用 Proxy 拦截对象读写，依赖收集与触发更新；比 Vue2 的 Object.defineProperty 更完整（可监听增删、数组等）。",
    ),
    (
        "ref 和 reactive 区别？",
        "ref 适合基本类型，取值需 .value；reactive 适合对象。模板中 ref 自动解包。",
    ),
    (
        "组件通信方式有哪些？",
        "props/emit、provide/inject、Pinia/Vuex、事件总线（少用）、路由参数、插槽。",
    ),
    (
        "什么是 SPA？有什么优缺点？",
        "单页应用：首屏加载后前端路由切换。体验流畅，但首屏可能慢，SEO 较弱（后台系统可接受）。",
    ),
    (
        "Vite 相对 Webpack 的优势？",
        "开发时用原生 ESM，冷启动与热更新更快；生产用 Rollup 打包。",
    ),
    (
        "如何优化前端性能？",
        "路由懒加载、图片压缩、防抖节流、减少重渲染、CDN、打包分析去冗余依赖、骨架屏。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.3 后端 / Java / Spring 面试题", 2)
hash_q = "#{" + "} 和 ${" + "} 区别？"
hash_a = (
    "#{" + "} 预编译防注入；${" + "} 字符串拼接有注入风险，仅用于动态表名/排序字段且必须白名单。"
)
for q, a in [
    (
        "Spring Boot 自动配置原理？",
        "通过 starter 引入依赖，条件注解按条件装配 Bean，约定优于配置，减少 XML。",
    ),
    (
        "Spring MVC 请求流程？",
        "DispatcherServlet 接收→HandlerMapping 找 Handler→拦截器 preHandle→Controller→返回→"
        "消息转换（JSON）→拦截器 afterCompletion。",
    ),
    (
        "你们 JWT 流程是怎样的？",
        "登录校验账号密码与审批状态→签发含 userId/username/role 的 token（如 7 天）→前端存储→"
        "请求头 Bearer 携带→JwtInterceptor 校验并写入 request 属性。",
    ),
    (
        "JWT 相对 Session 的优缺点？",
        "优点：无状态、易扩展、适合前后端分离。缺点：难主动失效（可用 Redis 黑名单）、token 被盗风险、载荷不宜过大。",
    ),
    (
        "为什么没用 Spring Security？",
        "教学/毕设阶段自定义拦截器更直观。生产可升级 Spring Security + JWT，统一认证、授权、密码编码等。",
    ),
    (
        "MyBatis 和 JPA/Hibernate 怎么选？",
        "复杂 SQL、联表、动态条件多时 MyBatis 更可控；JPA 开发快但复杂查询不够灵活。本项目选 MyBatis。",
    ),
    (hash_q, hash_a),
    (
        "PageHelper 原理简述？",
        "拦截 Executor，在 SQL 后拼分页语句并查 count，返回 Page 对象，再封装 total/list。",
    ),
    (
        "什么是事务？你们哪里需要事务？",
        "一组操作要么全成功要么全失败。如：审批通过同时写通知；可用 @Transactional。",
    ),
    (
        "IOC 和 AOP 是什么？",
        "IOC：控制反转，容器管理对象依赖注入。AOP：面向切面，如日志、事务、权限，不侵入业务代码。",
    ),
    (
        "常用注解有哪些？",
        "@RestController @RequestMapping @Autowired/@Resource @Service @Mapper @Transactional @ControllerAdvice 等。",
    ),
    (
        "GET 和 POST 区别？幂等性？",
        "GET 获取、可缓存、参数在 URL；POST 提交、写操作常在 Body。幂等：多次效果相同，PUT/DELETE 常幂等，POST 一般不幂等。",
    ),
    (
        "如何设计分页接口？",
        "入参 page、pageSize 与筛选条件；返回 total 与 list；避免一次拉全表。",
    ),
    (
        "全局异常怎么处理？",
        "@ControllerAdvice + @ExceptionHandler，如上传超限返回统一 Result.error。",
    ),
    (
        "连接池为什么重要？",
        "复用数据库连接，避免频繁创建销毁。Boot 默认 HikariCP，可配最大连接、最小空闲等。",
    ),
    (
        "HashMap 和 ConcurrentHashMap？",
        "HashMap 非线程安全；ConcurrentHashMap 保证并发安全，适合并发场景。",
    ),
    (
        "ArrayList 和 LinkedList？",
        "ArrayList 数组，随机访问快；LinkedList 链表，头尾插入删除相对方便。",
    ),
    (
        "== 和 equals？",
        "== 比较引用（基本类型比值）；equals 可重写比内容。字符串比较用 equals。",
    ),
    (
        "谈谈线程池",
        "复用线程、控制并发。核心参数：核心数、最大数、队列、拒绝策略。AI 异步任务可用线程池。",
    ),
    (
        "同步与异步、阻塞与非阻塞？",
        "同步等待结果，异步不等待；阻塞挂起线程，非阻塞立即返回。AI 评审适合异步，避免拖慢接口。",
    ),
    (
        "Bean 的作用域有哪些？",
        "常见 singleton（默认）、prototype；Web 下还有 request、session。",
    ),
    (
        "@Autowired 和 @Resource 区别？",
        "@Autowired 按类型，可配合 @Qualifier；@Resource 按名称优先（JSR-250）。",
    ),
    (
        "过滤器 Filter 和拦截器 Interceptor 区别？",
        "Filter 偏 Servlet 层；Interceptor 是 Spring MVC 机制，更易拿到 Handler 信息。项目 JWT 用拦截器。",
    ),
    (
        "如何防止表单重复提交？",
        "前端防抖、按钮置灰；后端 token 防重、唯一约束、状态机校验（已审批不可再批）。",
    ),
    (
        "RESTful 设计规范说几点？",
        "用名词资源路径；用 HTTP 方法表达动作；合理状态码；返回统一结构。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.4 MySQL 面试题", 2)
for q, a in [
    (
        "MySQL 索引有什么用？常见索引类型？",
        "加速查询。常见 B+Tree 索引、主键、唯一、联合索引；注意最左前缀。",
    ),
    (
        "什么情况下索引失效？",
        "对索引列函数/运算、隐式类型转换、前导模糊 %like、OR 不当、优化器判断全表更快等。",
    ),
    ("事务 ACID？", "原子性、一致性、隔离性、持久性。"),
    (
        "隔离级别与脏读/不可重复读/幻读？",
        "读未提交/读已提交/可重复读/串行化。MySQL InnoDB 默认可重复读，结合 MVCC 与锁机制。",
    ),
    (
        "InnoDB 和 MyISAM？",
        "InnoDB 支持事务、行锁、外键；MyISAM 不支持事务，表锁。现多用 InnoDB。",
    ),
    (
        "主键、外键、唯一键？",
        "主键唯一非空；外键维护关联（本项目部分表 ON DELETE CASCADE）；唯一键保证业务唯一如学号。",
    ),
    (
        "如何优化慢 SQL？",
        "explain 看执行计划、加合适索引、避免 select *、减少大偏移分页、必要冗余/缓存。",
    ),
    ("varchar 和 char？", "char 定长，varchar 变长更省空间。姓名手机等多用 varchar。"),
    (
        "为什么用 utf8mb4？",
        "完整支持 emoji 与更多 Unicode 字符，比 utf8（3字节）更完整。",
    ),
    (
        "什么是最左前缀原则？",
        "联合索引 (a,b,c) 可用 a、a+b、a+b+c；缺 a 时一般用不了该联合索引。",
    ),
    (
        "count(*) / count(1) / count(列)？",
        "InnoDB 中 count(*) 与 count(1) 通常相当；count(列) 不计 NULL。",
    ),
    (
        "什么是覆盖索引？",
        "查询列都在索引中，无需回表，性能更好。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.5 Redis 面试题", 2)
for q, a in [
    (
        "为什么要用 Redis？",
        "内存读写快，适合热点数据缓存、会话、计数、限流，减轻数据库压力、降低接口延迟。",
    ),
    (
        "Redis 为什么快？",
        "内存操作、单线程避免上下文切换（命令执行）、高效数据结构、I/O 多路复用。",
    ),
    (
        "缓存穿透、击穿、雪崩？如何解决？",
        "穿透：不存在的 key 打 DB→空值缓存/布隆。击穿：热点过期→互斥锁/逻辑过期。"
        "雪崩：同时过期→随机 TTL、集群。",
    ),
    (
        "Redis 持久化？",
        "RDB 快照恢复快、可能丢数据；AOF 更安全、文件大；可混合。",
    ),
    (
        "Redis 分布式锁怎么做？",
        "SET key value NX EX；释放用 Lua 校验 value 再删。注意误删与续期（Redisson 看门狗）。",
    ),
    (
        "热 Key、大 Key 问题？",
        "热 Key 可本地缓存/拆分；大 Key 拆分存储、避免大 value 阻塞。",
    ),
    (
        "如何在本项目用 Redis？",
        "验证码 TTL、未读数、JWT 黑名单、AI 进度、接口限流。",
    ),
    (
        "Redis 和 Memcached 区别？",
        "Redis 数据结构更丰富、可持久化；Memcached 更单纯的内存缓存。",
    ),
    (
        "什么是缓存预热？",
        "系统启动或低峰时提前把热点数据载入缓存，避免刚上线就打爆 DB。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.6 微服务 / 消息队列 / 部署", 2)
for q, a in [
    (
        "什么是微服务？优缺点？",
        "按业务拆独立可部署服务。优点：独立扩容、团队并行。缺点：分布式事务、运维、调用链路复杂。",
    ),
    (
        "服务注册发现是什么？",
        "服务启动向注册中心登记，调用方从注册中心获取实例列表，实现动态上下线。",
    ),
    (
        "网关的作用？",
        "统一入口：路由、鉴权、限流、日志，避免每个服务重复实现。",
    ),
    (
        "什么是熔断降级？",
        "下游故障时快速失败或返回兜底，防止雪崩；非核心功能可降级。",
    ),
    (
        "分布式事务怎么解决？",
        "尽量避免强一致；用最终一致性：本地消息表、MQ 事务消息、Saga、Seata 等。",
    ),
    (
        "消息队列解决什么问题？",
        "异步解耦、削峰填谷、最终一致性。如审批后发通知、AI 任务投递。",
    ),
    (
        "RabbitMQ 和 Kafka 区别（简）？",
        "RabbitMQ 路由灵活偏业务消息；Kafka 高吞吐偏日志流、大数据。",
    ),
    (
        "Docker 是什么？",
        "容器化打包应用与依赖，环境一致，便于部署。可用 docker-compose 编排 MySQL、后端、Nginx。",
    ),
    (
        "Nginx 反向代理和负载均衡？",
        "反向代理把请求转到后端；负载均衡把请求分到多实例（轮询、加权等）。",
    ),
    (
        "如何保证接口幂等？",
        "唯一业务单号、token 防重、数据库唯一约束、状态机校验（已审批不可再批）。",
    ),
    (
        "什么是服务雪崩？",
        "一个下游故障引起连锁超时，线程耗尽，整个系统不可用。用超时、隔离舱、熔断限流应对。",
    ),
    (
        "配置中心有什么用？",
        "集中管理配置，支持动态刷新，避免每个服务改配置重新打包。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.7 计算机网络 / HTTP / 安全", 2)
for q, a in [
    (
        "HTTP 和 HTTPS？",
        "HTTPS = HTTP + TLS，加密传输，防窃听篡改，需证书。生产必须 HTTPS。",
    ),
    (
        "HTTP 常见状态码？",
        "200 成功；301/302 重定向；400 参数错误；401 未认证；403 无权限；404 不存在；500 服务器错误。",
    ),
    (
        "一次完整的 URL 请求过程？",
        "DNS→建立连接（TCP/TLS）→发 HTTP 请求→服务器处理→响应→浏览器渲染；可能有缓存与 CDN。",
    ),
    (
        "TCP 和 UDP？",
        "TCP 可靠面向连接；UDP 快但不保证可靠。HTTP 基于 TCP。",
    ),
    (
        "什么是 XSS、CSRF？",
        "XSS：注入脚本偷信息→转义/CSP。CSRF：借用户身份发请求→Token/SameSite Cookie。"
        "JWT 放 Header 相对利于防部分 CSRF。",
    ),
    (
        "SQL 注入如何防？",
        "预编译参数（MyBatis 井号占位）、输入校验、最小权限账号。",
    ),
    (
        "什么是同源策略？",
        "协议+域名+端口都相同才算同源，限制恶意页面读写其他源数据。",
    ),
    (
        "Cookie / Session / Token 区别？",
        "Cookie 存浏览器；Session 存服务端；Token（如 JWT）常存客户端并由服务端校验签名。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.8 操作系统 / 场景设计 / 行为面", 2)
for q, a in [
    (
        "进程和线程区别？",
        "进程是资源分配单位，线程是调度单位；同进程线程共享内存，通信更方便也更需同步。",
    ),
    (
        "什么是死锁？四个条件？",
        "互相等待对方资源。互斥、占有且等待、不可抢占、循环等待。破坏其一即可。",
    ),
    (
        "设计一个登录功能要注意什么？",
        "密码加密、错误次数限制、验证码、JWT/Session、HTTPS、状态校验（是否审批）、日志审计。",
    ),
    (
        "设计通知系统？",
        "写通知表；未读计数可用 Redis；推送可用轮询/WebSocket；分类 type；已读标记。",
    ),
    (
        "你的优缺点？",
        "优点结合项目：能闭环交付、肯钻研 AI 接入。缺点选真实可改进的，如测试覆盖不足、尚未上 Redis，并说明学习计划。",
    ),
    (
        "职业规划？",
        "先扎实做好业务开发与基础（Java/Vue/MySQL），再补 Redis、MQ、微服务，逐步能独立负责模块。",
    ),
    (
        "为什么选我们公司？",
        "匹配技术栈、成长空间、业务方向；不要贬低其他公司。结合 JD 说 1～2 个吸引点。",
    ),
    (
        "有什么要问我的？",
        "可问：技术栈与团队规模、培养机制、业务方向、试用期考核。体现真诚与思考。",
    ),
    (
        "项目中最大的收获？",
        "完整走通需求分析→库表→接口→页面→联调部署；理解权限与状态机；实践 AI 接口工程化。",
    ),
    (
        "如果重做一遍，你会改什么？",
        "密码改 BCrypt；引入 Redis；文件上 OSS；通知/AI 用 MQ 或 WebSocket；补接口文档与单测。",
    ),
]:
    add_qa(q, a)

add_heading_cn("5.9 快速追问（短答卡）", 2)
for s in [
    "IOC：控制反转，对象交给容器管理。",
    "AOF/RDB：Redis 两种持久化方式。",
    "幂等：同一请求执行多次结果一致。",
    "Base64 不是加密，是编码。",
    "JWT 三段：Header.Payload.Signature。",
    "REST：用 HTTP 方法表达资源操作。",
    "前后端分离靠 JSON API 通信。",
    "乐观锁/悲观锁：版本号 vs 真正加锁。",
    "CDN：就近缓存加速静态资源。",
    "CI/CD：持续集成与持续部署。",
    "MVCC：多版本并发控制，提升读性能。",
    "B+树：MySQL InnoDB 索引底层结构。",
    "CAP：一致性、可用性、分区容忍（分布式取舍）。",
    "三次握手 / 四次挥手：TCP 建连与断开。",
]:
    add_para("• " + s, size=10.5)

# ========== 六七 ==========
add_heading_cn("六、面试临场策略", 1)
for i in [
    "1. 先用 1.5 分钟自我介绍稳住节奏，再等面试官深挖项目。",
    "2. 每个难点按「背景→问题→方案→结果→可改进」说，避免只罗列技术名词。",
    "3. 不会的题：诚实说了解有限，再讲思路或与项目关联的改进方向。",
    "4. 主动提不足：MD5→BCrypt、本地文件→OSS、轮询→WebSocket、引入 Redis/MQ。",
    "5. 准备一次完整演示话术：登录→申报→上传计划书→审批→通知→AI 评审。",
]:
    add_para(i)

add_heading_cn("七、建议背诵优先级", 1)
add_para("P0：自我介绍、项目状态流转、JWT 流程、AI 评审链路、统一响应与权限。")
add_para("P1：Vue3/Pinia/Axios、MyBatis 预编译、事务、索引、Redis 三大缓存问题。")
add_para("P2：微服务组件、MQ、Docker/Nginx、HTTPS/状态码。")
add_para("P3：操作系统、集合源码、手写算法（按岗位另加）。")

end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
re = end.add_run("—— 文档结束，祝面试顺利 ——")
set_run_font(re, size=11, bold=True, color=(0x1A, 0x56, 0xDB))

out = r"c:\Users\Lenovo\Desktop\智能创新创业项目管理系统\code2\面试准备手册-创新创业项目管理系统.docx"
doc.save(out)
print("OK:", out)
